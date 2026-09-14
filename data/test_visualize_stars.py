import csv
from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

import visualize_stars as v


class ViewerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.folder = Path(self.temp.name)

    def csv_file(self, rows):
        path = self.folder / "stars.csv"
        with path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.DictWriter(f, fieldnames=rows[0].keys())
            writer.writeheader()
            writer.writerows(rows)
        return path

    def row(self, **changes):
        return {"hyg_id": "1", "star_name": "Test star", "distance_ly": "5",
                "x_ly": "3", "y_ly": "4", "z_ly": "0", "planet_count": "2",
                "planet_count_override": "", **changes}

    def test_csv_coordinates_and_missing_sol_reference(self):
        stars = v.load_stars(self.csv_file([self.row()]))
        self.assertTrue(stars[0]["reference_only"])
        self.assertEqual([stars[0][a] for a in "xyz"], [0, 0, 0])
        self.assertEqual([stars[1][a] for a in "xyz"], [3, 4, 0])

    def test_sol_is_exactly_origin_without_translation_of_other_stars(self):
        stars = v.load_stars(self.csv_file([self.row(hyg_id="0", star_name="Sol", distance_ly="0", x_ly="0.00001"), self.row()]))
        self.assertEqual(len(stars), 2)
        self.assertEqual(stars[0]["x"], 0)
        self.assertEqual(stars[1]["x"], 3)

    def test_csv_human_count_override(self):
        stars = v.load_stars(self.csv_file([self.row(planet_count_override="0")]))
        self.assertEqual(stars[1]["planet_count"], 0)
        self.assertTrue(stars[1]["count_edited"])

    def test_workbook_formulas_without_caches_and_planet_exclusions(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "Stars"
        row = self.row(planet_count='=IF(D2="",E2+0,D2)')
        ws.append(list(row))
        ws.append(list(row.values()))
        planets = wb.create_sheet("Planets")
        planets.append(["hyg_id", "include_in_count"])
        planets.append([1, 1])
        planets.append([1, 0])
        path = self.folder / "edited.xlsx"
        wb.save(path)
        self.assertEqual(v.load_stars(path)[1]["planet_count"], 1)
        ws.cell(2, list(row).index("planet_count_override") + 1, 7)
        wb.save(path)
        self.assertEqual(v.load_stars(path)[1]["planet_count"], 7)

    def test_bad_coordinates_counts_and_duplicate_ids_fail(self):
        for row in [self.row(x_ly="nan"), self.row(z_ly=""), self.row(planet_count="2.5"), self.row(planet_count="-1")]:
            with self.subTest(row=row), self.assertRaises(ValueError):
                v.load_stars(self.csv_file([row]))
        with self.assertRaises(ValueError):
            v.load_stars(self.csv_file([self.row(), self.row()]))

    def test_html_is_offline_and_source_text_cannot_close_script(self):
        source = self.csv_file([self.row(star_name='</script><script>alert(1)</script>')])
        stars = v.load_stars(source)
        output = self.folder / "map.html"
        v.write_viewer(stars, source, output)
        document = output.read_text()
        self.assertNotIn('<script src=', document)
        self.assertNotIn('</script><script>alert(1)', document)
        self.assertIn('\\u003c/script>', document)
        figure, _ = v.build_figure(stars)
        self.assertEqual(list(figure.data[0].x), [3])
        self.assertEqual(list(figure.data[1].x), [0])


if __name__ == "__main__":
    unittest.main()
