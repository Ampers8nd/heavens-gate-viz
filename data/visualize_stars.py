#!/usr/bin/env python3
"""Build an offline, interactive 3D star map from the selected XLSX or CSV."""

import argparse
from collections import Counter
import csv
import html
import json
import math
from pathlib import Path
import sys
import webbrowser

ROOT = Path(__file__).resolve().parent
DEFAULT_INPUT = ROOT / "output/nearby_30ly/nearby_stars.xlsx"


def finite(value, field):
    try:
        result = float(value)
    except (ValueError, TypeError) as exc:
        raise ValueError(f"{field}: expected a number, got {value!r}") from exc
    if not math.isfinite(result):
        raise ValueError(f"{field}: must be finite")
    return result


def count(value, field):
    value = finite(value, field)
    if value < 0 or not value.is_integer():
        raise ValueError(f"{field}: expected a nonnegative integer")
    return int(value)


def identifier(value):
    if value is None:
        return ""
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def worksheet_rows(sheet):
    rows = sheet.iter_rows(values_only=True)
    columns = next(rows, ())
    return [dict(zip(columns, row)) for row in rows if any(v is not None for v in row)]


def load_rows(path):
    """Recompute this project's known workbook formulas, without relying on caches."""
    if path.suffix.lower() == ".csv":
        with path.open(encoding="utf-8-sig", newline="") as f:
            return list(csv.DictReader(f))
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Input must be a selected stars .xlsx workbook or .csv file")
    from openpyxl import load_workbook
    with path.open("rb") as handle:
        wb = load_workbook(handle, read_only=True, data_only=False)
        try:
            if "Stars" not in wb:
                raise ValueError("Workbook must contain a Stars sheet")
            rows = worksheet_rows(wb["Stars"])
            included = Counter()
            if "Planets" in wb:
                for p in worksheet_rows(wb["Planets"]):
                    flag = p.get("include_in_count")
                    if flag in (None, ""):
                        continue
                    flag = count(flag, "Planets include_in_count")
                    if flag not in (0, 1):
                        raise ValueError("Planets include_in_count must be 0 or 1")
                    if flag:
                        included[identifier(p.get("hyg_id"))] += 1
            for row in rows:
                value = row.get("planet_count")
                override = row.get("planet_count_override")
                if override not in (None, ""):
                    row["planet_count"] = count(override, "planet_count_override")
                elif isinstance(value, str) and value.startswith("="):
                    if not value.startswith("=IF(D") or "Planets" not in wb:
                        raise ValueError("Unrecognized count formula; use planet_count_override or export calculated CSV")
                    solar = row.get("solar_system_planet_count") or 0
                    row["planet_count"] = included[identifier(row.get("hyg_id"))] + count(solar, "solar_system_planet_count")
            return rows
        finally:
            wb.close()


def load_stars(path):
    stars, seen = [], set()
    for row in load_rows(path):
        star_id = identifier(row.get("hyg_id"))
        if not star_id or star_id in seen:
            raise ValueError(f"Missing or duplicate HYG ID: {star_id!r}")
        seen.add(star_id)
        name = str(row.get("star_name") or row.get("proper_name") or row.get("gliese") or f"HYG {star_id}")
        sol = name.casefold() in {"sol", "sun"} and finite(row.get("distance_ly"), "distance_ly") == 0
        coords = [0.0, 0.0, 0.0] if sol else [finite(row.get(f"{axis}_ly"), f"{name} {axis}_ly") for axis in "xyz"]
        override = row.get("planet_count_override")
        value = override if override not in (None, "") else row.get("planet_count")
        planets = count(value, f"{name} planet_count")
        distance = finite(row.get("distance_ly"), f"{name} distance_ly")
        if distance < 0:
            raise ValueError(f"{name}: distance_ly cannot be negative")
        aliases = [str(row.get(k) or "") for k in ("proper_name", "gliese", "bayer_flamsteed", "nasa_hosts")]
        aliases += [f"{k.upper()} {row[k]}" for k in ("hip", "hd", "hr") if row.get(k)]
        stars.append({"id": star_id, "name": name, "x": coords[0], "y": coords[1], "z": coords[2],
                      "distance": distance, "planet_count": planets, "sol": sol,
                      "planet_names": str(row.get("planet_names") or ""),
                      "notes": str(row.get("notes") or ""), "spectral_type": str(row.get("spectral_type") or ""),
                      "aliases": "; ".join(dict.fromkeys(a for a in aliases if a)),
                      "count_edited": override not in (None, ""), "reference_only": False})
    if sum(s["sol"] for s in stars) > 1:
        raise ValueError("Multiple Sol entries")
    # Keep the Sun as a visual reference even for a spreadsheet exported with --exclude-sol.
    if not any(s["sol"] for s in stars):
        stars.insert(0, {"id": "sol-reference", "name": "Sol", "x": 0, "y": 0, "z": 0,
                         "distance": 0, "planet_count": 8, "sol": True, "planet_names": "",
                         "notes": "Origin reference added by viewer; absent from the selected spreadsheet.",
                         "spectral_type": "G2V", "aliases": "Sun", "count_edited": False, "reference_only": True})
    return stars


def build_figure(stars):
    import plotly.graph_objects as go

    extent = max(5, max(math.sqrt(sum(s[a] ** 2 for a in "xyz")) for s in stars) * 1.08)
    field = [s for s in stars if not s["sol"]]
    sun = next(s for s in stars if s["sol"])
    figure = go.Figure()
    figure.add_trace(go.Scatter3d(
        x=[s["x"] for s in field], y=[s["y"] for s in field], z=[s["z"] for s in field],
        customdata=[s["id"] for s in field], mode="markers", text=[""] * len(field),
        textposition="top center", textfont=dict(size=11, color="#ccdeef"),
        marker=dict(size=[4 if s["planet_count"] else 2.6 for s in field],
                    color=["#67ded6" if s["planet_count"] else "#97aabe" for s in field], opacity=1),
        hovertext=[html.escape(s["name"]) for s in field], hovertemplate="%{hovertext}<extra></extra>",
        name="Stars", showlegend=False))
    figure.add_trace(go.Scatter3d(x=[0], y=[0], z=[0], customdata=[sun["id"]],
        mode="markers", text=["SOL"], textposition="top center",
        textfont=dict(size=18, color="#ffdc87"), marker=dict(size=11, color="#ffd477"),
        hovertemplate="Sol · origin<extra></extra>", name="Sol", showlegend=False))
    figure.add_trace(go.Scatter3d(x=[], y=[], z=[], mode="markers", text=[], customdata=[],
        textposition="top center", textfont=dict(size=15, color="white"),
        marker=dict(size=7, color="#ffffff", symbol="diamond-open"),
        hoverinfo="skip", name="Selected star", showlegend=False))
    # Faint reference rings, not fictitious stars or orbital paths.
    ring_x, ring_y, ring_z = [], [], []
    for radius in range(10, math.ceil(extent), 10):
        for step in range(129):
            angle = 2 * math.pi * step / 128
            ring_x.append(radius * math.cos(angle))
            ring_y.append(radius * math.sin(angle))
            ring_z.append(0)
        ring_x.append(None); ring_y.append(None); ring_z.append(None)
    figure.add_trace(go.Scatter3d(x=ring_x, y=ring_y, z=ring_z, mode="lines",
        line=dict(color="#253347", width=1), hoverinfo="skip", showlegend=False))
    axis = dict(range=[-extent, extent], backgroundcolor="#080f1c", showbackground=False,
                gridcolor="#172638", zerolinecolor="#304354", showspikes=False,
                tickfont=dict(size=10, color="#6d8399"), nticks=5)
    figure.update_layout(
        paper_bgcolor="#080f1c", plot_bgcolor="#080f1c", font=dict(color="#b6c9dc"),
        margin=dict(l=0, r=0, t=0, b=0), hovermode="closest", uirevision="star-map",
        scene=dict(aspectmode="cube", dragmode="orbit", bgcolor="#080f1c",
                   xaxis={**axis, "title": "X · ly"}, yaxis={**axis, "title": "Y · ly"},
                   zaxis={**axis, "title": "Z · ly"},
                   camera=dict(eye=dict(x=1.5, y=1.5, z=1.1), center=dict(x=0, y=0, z=0),
                               up=dict(x=0, y=0, z=1))))
    return figure, extent


def write_viewer(stars, source, output):
    import plotly.io as pio

    figure, extent = build_figure(stars)
    template = (ROOT / "viewer.html").read_text(encoding="utf-8")
    payload = json.dumps({"stars": stars, "extent": extent, "source": source.name}, ensure_ascii=False)
    payload = payload.replace("<", "\\u003c").replace("\u2028", "\\u2028").replace("\u2029", "\\u2029")
    chart = pio.to_html(figure, full_html=False, include_plotlyjs=True, div_id="star-map",
                       config={"scrollZoom": True, "displayModeBar": False, "responsive": True},
                       post_script="window.dispatchEvent(new Event('star-map-ready'));")
    document = template.replace("<!-- STAR_DATA -->", f'<script id="star-data" type="application/json">{payload}</script>')
    document = document.replace("<!-- PLOT -->", chart)
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(document, encoding="utf-8")


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT, help="Selected XLSX (Stars/Planets sheets) or stars.csv")
    parser.add_argument("--output", type=Path, help="HTML output; defaults to star_map.html beside input")
    parser.add_argument("--open", action="store_true", help="Open the generated map in your default browser")
    args = parser.parse_args(argv)
    try:
        stars = load_stars(args.input)
        output = args.output or args.input.parent / "star_map.html"
        if output.resolve() == args.input.resolve():
            raise ValueError("Output must not overwrite the input spreadsheet")
        if output.suffix.lower() != ".html":
            raise ValueError("Output filename must end in .html")
        write_viewer(stars, args.input, output)
        print(f"Rendered {sum(not s['reference_only'] for s in stars)} spreadsheet entries, with Sol at (0, 0, 0).")
        print(f"Open in a browser: {output.resolve()}")
        if args.open and not webbrowser.open(output.resolve().as_uri()):
            print("No browser could be opened automatically; open the HTML file manually.")
        return 0
    except (ValueError, OSError, ImportError, csv.Error) as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    sys.exit(main())
