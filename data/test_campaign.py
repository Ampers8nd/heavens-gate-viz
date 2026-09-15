from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook

from campaign import apply_campaign


class CampaignPlanetTests(unittest.TestCase):
    def make_workbook(self, path):
        workbook = Workbook()
        stars = workbook.active
        stars.title = "Stars"
        stars.append(["hyg_id", "system_id", "faction_owners", "planet_count_override", "notes"])
        stars.append([1, "alpha-system", "union", None, "Campaign star"])
        planets = workbook.create_sheet("Planets")
        planets.append(["planet_name", "host_name", "hyg_id", "include_in_count", "notes", "faction_owners"])
        planets.append(["Archive b", "Alpha", 1, 1, "Archive notes", None])
        planets.append(["Campaign Prime", "Alpha", 1, 1, "Artificial world", None])
        planets.append(["Outside", "Missing", 999, 1, "Outside selected stars", None])
        factions = workbook.create_sheet("Factions")
        factions.append(["faction_id", "faction_name", "color", "notes"])
        factions.append(["union", "The Union", "#112233", None])
        workbook.save(path)
        workbook.close()

    def test_preserves_manual_planet_for_selected_host(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "campaign.xlsx"
            self.make_workbook(source)
            stars = [{"hyg_id": "1", "primary_hyg_id": "1", "notes": ""}]
            planets = [{"planet_name": "Archive b", "host_name": "Alpha", "hyg_id": "1",
                        "include_in_count": 1, "notes": "", "faction_owners": ""}]

            apply_campaign(stars, planets, source)

            self.assertEqual([planet["planet_name"] for planet in planets], ["Archive b", "Campaign Prime"])
            self.assertEqual(planets[1]["notes"], "Artificial world")

    def test_does_not_duplicate_existing_planet(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "campaign.xlsx"
            self.make_workbook(source)
            stars = [{"hyg_id": "1", "primary_hyg_id": "1", "notes": ""}]
            planets = [{"planet_name": "Campaign Prime", "host_name": "Alpha", "hyg_id": "1",
                        "include_in_count": 1, "notes": "", "faction_owners": ""}]

            apply_campaign(stars, planets, source)

            self.assertEqual(sum(p["planet_name"] == "Campaign Prime" for p in planets), 1)


if __name__ == "__main__":
    unittest.main()
