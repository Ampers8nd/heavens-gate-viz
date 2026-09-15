#!/usr/bin/env python3
"""Join a local HYG catalog to a NASA Planetary Systems export, offline."""

import argparse
from collections import defaultdict
import csv
import gzip
import hashlib
import html
import json
import math
from pathlib import Path
import re
import sys
import unicodedata
from campaign import FACTION_COLUMNS, OWNERSHIP_HELP, apply_campaign

PC_TO_LY = 3.2615637771674333
DEFAULT_RADIUS_LY = 30.0
ROOT = Path(__file__).resolve().parent

STAR_COLUMNS = [
    "star_name", "has_planets", "planet_count", "planet_count_override",
    "confirmed_exoplanet_count", "distance_ly", "planet_names", "notes",
    "hyg_id", "nasa_hosts", "match_method", "controversial_planet_count",
    "distance_pc", "spectral_type", "hyg_luminosity_solar", "apparent_magnitude",
    "hip", "hd", "hr", "gliese", "bayer_flamsteed", "proper_name",
    "ra_deg_j2000", "dec_deg_j2000", "x_ly", "y_ly", "z_ly",
    "component", "primary_hyg_id", "archive_planet_count", "solar_system_planet_count",
    "x_pc", "y_pc", "z_pc",
    "system_id", "faction_owners",
]
PARAMETERS = {
    "semi_major_axis_au": "pl_orbsmax", "orbital_period_days": "pl_orbper",
    "eccentricity": "pl_orbeccen", "insolation_earth": "pl_insol",
    "equilibrium_temperature_k": "pl_eqt", "radius_earth": "pl_rade",
    "mass_or_msini_earth": "pl_bmasse", "stellar_temperature_k": "st_teff",
    "stellar_mass_solar": "st_mass", "stellar_radius_solar": "st_rad",
    "stellar_metallicity_dex": "st_met",
}
PLANET_COLUMNS = [
    "planet_name", "host_name", "hyg_id", "hyg_star_name", "include_in_count", "notes",
    *PARAMETERS, "mass_provenance", "controversial", "discovery_method", "discovery_year",
    "nasa_distance_ly", "hyg_distance_ly", "distance_difference_percent", "match_method",
    "angular_separation_arcsec", "parameter_selection", "source_rows",
    "system_planet_count", "system_star_count", "solution_type", "reference",
    "reference_url", "stellar_reference", "system_reference", "row_updated",
    "faction_owners",
]
AUDIT_COLUMNS = [
    "nasa_host", "planet_count", "status", "hyg_id", "hyg_name", "hyg_distance_ly",
    "nasa_distance_ly", "method", "evidence", "angular_separation_arcsec",
    "distance_difference_percent", "candidate_hyg_ids", "notes",
]


def metadata_filename(path):
    """Return a portable source label without exposing a local directory."""
    return Path(path).name if path else None


def number(value):
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def scaled(value, factor=1):
    n = number(value)
    return n * factor if n is not None else ""


def normalize(name):
    """Normalize spelling/spacing, retaining catalog signs and component suffixes."""
    name = unicodedata.normalize("NFKC", name).strip().casefold().replace("−", "-")
    name = re.sub(r"^(?:gliese|gj|gl)\s*(?=\d)", "gj", name)
    name = re.sub(r"\s+", "", name).replace("'", "").replace("’", "")
    # Bayer superscripts are represented as e.g. Alp-2 in HYG.
    return re.sub(r"^([a-z]{3})-(\d)", r"\1\2", name)


def identifiers(star):
    names = [star.get(k, "") for k in ("gl", "proper", "bf")]
    names += [f"{k.upper()} {star[k]}" for k in ("hip", "hd", "hr") if star.get(k)]
    names += [f"HYG {star['id']}"]
    for key in ("bayer", "flam", "var"):
        if star.get(key) and star.get("con"):
            names.append(f"{star[key]} {star['con']}")
    return {normalize(n) for n in names if n}


def star_name(star):
    return (star.get("proper") or star.get("gl") or star.get("bf")
            or (f"HIP {star['hip']}" if star.get("hip") else f"HYG {star['id']}"))


def is_sol(star):
    return star.get("proper", "").casefold() in {"sol", "sun"} and number(star["dist"]) == 0


def cartesian(star, factor=1):
    """HYG's heliocentric J2000 equatorial axes; pin Sol exactly to the origin."""
    return {axis: 0.0 if is_sol(star) else scaled(star.get(axis), factor) for axis in "xyz"}


def selected(star, radius, exclude_sol=False):
    dist = number(star.get("dist"))
    return (dist is not None and 0 <= dist < 100000 and dist * PC_TO_LY <= radius
            and not (exclude_sol and is_sol(star)))


def read_csv(path, required):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(line for line in handle if not line.startswith("#"))
        missing = set(required) - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path.name}: missing columns: {', '.join(sorted(missing))}")
        return list(reader), reader.fieldnames


def choose_planets(rows):
    """PS membership defines the confirmed list; soltype describes a parameter set."""
    grouped = defaultdict(list)
    for row in rows:
        if not row["pl_name"].strip() or not row["hostname"].strip():
            raise ValueError("NASA rows must have a planet name and host name")
        grouped[row["pl_name"]].append(row)
    chosen = {}
    for name, group in grouped.items():
        if len({r["hostname"] for r in group}) != 1:
            raise ValueError(f"Conflicting host names for planet {name}")
        defaults = [r for r in group if r.get("default_flag") == "1"]
        pool = defaults or group
        # Deterministic date tie-break; preserve a single self-consistent source row.
        row = max(pool, key=lambda r: (r.get("rowupdate", ""), r.get("pl_pubdate", ""),
                                       json.dumps(r, sort_keys=True)))
        policy = "NASA default" if len(defaults) == 1 else (
            "latest of multiple defaults" if defaults else "latest row; no default provided")
        chosen[name] = (row, policy)
    return chosen, grouped


def separation(host, star):
    values = [number(host.get("ra")), number(host.get("dec")),
              number(star.get("ra")), number(star.get("dec"))]
    if any(v is None for v in values):
        return None
    ra, dec, sra, sdec = values
    ra, dec, sra, sdec = map(math.radians, (ra, dec, sra * 15, sdec))
    h = math.sin((dec - sdec) / 2) ** 2 + math.cos(dec) * math.cos(sdec) * math.sin((ra - sra) / 2) ** 2
    return math.degrees(2 * math.asin(math.sqrt(max(0, min(1, h))))) * 3600


def distance_difference(host, star):
    nasa, hyg = number(host.get("sy_dist")), number(star.get("dist"))
    return 100 * abs(nasa - hyg) / hyg if nasa is not None and hyg else ""


def load_aliases(path, index, hosts):
    aliases = {}
    if path is None:
        return aliases
    rows, _ = read_csv(path, ["nasa_hostname", "hyg_identifier", "source", "notes"])
    for row in rows:
        name = row["nasa_hostname"]
        if name in aliases:
            raise ValueError(f"Duplicate alias for {name}")
        hits = index.get(normalize(row["hyg_identifier"]), set())
        if name in hosts and len(hits) != 1:
            raise ValueError(f"Alias for {name} must identify exactly one HYG entry: {hits}")
        aliases[name] = (hits, row)
    return aliases


def match_hosts(stars, chosen, radius, alias_path, review_arcsec, exclude_sol):
    by_id = {s["id"]: s for s in stars}
    if len(by_id) != len(stars):
        raise ValueError("HYG contains duplicate IDs")
    index = defaultdict(set)
    for s in stars:
        if not is_sol(s):
            for name in identifiers(s):
                index[name].add(s["id"])
    hosts = defaultdict(list)
    for row, _ in chosen.values():
        hosts[row["hostname"]].append(row)
    aliases = load_aliases(alias_path, index, hosts)
    nearby = [s for s in stars if selected(s, radius, exclude_sol) and not is_sol(s)]
    matches, audit = {}, []
    for host, planets in sorted(hosts.items()):
        representative = sorted(planets, key=lambda r: r["pl_name"])[0]
        hits = index.get(normalize(host), set())
        method, evidence = "catalog identifier", host
        if host in aliases:
            hits, alias = aliases[host]
            method = "verified alias"
            evidence = f"{alias['hyg_identifier']}; {alias['source']}; {alias['notes']}"
        entry = dict.fromkeys(AUDIT_COLUMNS, "")
        entry.update(nasa_host=host, planet_count=len(planets), method=method, evidence=evidence,
                     nasa_distance_ly=scaled(representative.get("sy_dist"), PC_TO_LY))
        if len(hits) == 1:
            s = by_id[next(iter(hits))]
            within = selected(s, radius, exclude_sol)
            diff = distance_difference(representative, s)
            entry.update(status="matched_selected" if within else "matched_outside_radius",
                         hyg_id=s["id"], hyg_name=star_name(s), hyg_distance_ly=scaled(s["dist"], PC_TO_LY),
                         angular_separation_arcsec=separation(representative, s), distance_difference_percent=diff)
            if diff != "" and diff > 20:
                entry["notes"] = "Catalog distances differ by >20%; HYG controls selection."
            if within:
                matches[host] = (s, method)
        else:
            candidates = set(hits)
            for s in nearby:
                angle = separation(representative, s)
                if angle is not None and angle <= review_arcsec:
                    candidates.add(s["id"])
            nasa_near = any(0 < (number(p.get("sy_dist")) or math.inf) * PC_TO_LY <= radius for p in planets)
            entry.update(status="ambiguous_identifier" if hits else (
                "unmatched_nearby" if nasa_near or candidates else "unmatched_other"),
                candidate_hyg_ids="; ".join(sorted(candidates, key=int)),
                method="unassigned", notes="Position candidates require review; component identity is preserved.")
        audit.append(entry)
    return matches, audit


def plain_reference(value):
    return html.unescape(re.sub(r"<[^>]+>", "", value)).strip()


def reference_url(value):
    match = re.search(r"href\s*=\s*[\"']?([^\s>\"']+)", value)
    return html.unescape(match[1]) if match else ""


def make_tables(stars, chosen, grouped, matches, radius, exclude_sol):
    planets, measurements = [], []
    by_star = defaultdict(list)
    for name, (row, policy) in sorted(chosen.items()):
        if row["hostname"] not in matches:
            continue
        s, method = matches[row["hostname"]]
        p = dict.fromkeys(PLANET_COLUMNS, "")
        p.update(planet_name=name, host_name=row["hostname"], hyg_id=s["id"],
                 hyg_star_name=star_name(s), include_in_count=1,
                 **{label: scaled(row.get(field)) for label, field in PARAMETERS.items()},
                 mass_provenance=row.get("pl_bmassprov", ""), controversial=int(row.get("pl_controv_flag") or 0),
                 discovery_method=row.get("discoverymethod", ""), discovery_year=scaled(row.get("disc_year")),
                 nasa_distance_ly=scaled(row.get("sy_dist"), PC_TO_LY), hyg_distance_ly=scaled(s["dist"], PC_TO_LY),
                 distance_difference_percent=distance_difference(row, s), match_method=method,
                 angular_separation_arcsec=separation(row, s), parameter_selection=policy,
                 source_rows=len(grouped[name]), system_planet_count=scaled(row.get("sy_pnum")),
                 system_star_count=scaled(row.get("sy_snum")), solution_type=row.get("soltype", ""),
                 reference=plain_reference(row.get("pl_refname", "")), reference_url=reference_url(row.get("pl_refname", "")),
                 stellar_reference=plain_reference(row.get("st_refname", "")),
                 system_reference=plain_reference(row.get("sy_refname", "")), row_updated=row.get("rowupdate", ""))
        limits = [f"{label}: {'upper' if number(row[field+'lim']) == 1 else 'lower'} limit"
                  for label, field in PARAMETERS.items() if number(row.get(field + "lim")) in (1, -1)]
        p["notes"] = "; ".join(limits)
        planets.append(p)
        by_star[s["id"]].append(p)
        for source_row in grouped[name]:
            # Keep source identifiers/text literal; expose numerical measurements as numbers.
            values = {k: (number(v) if k not in {"pl_name", "hostname"} and number(v) is not None else v)
                      for k, v in source_row.items()}
            measurements.append({"hyg_id": s["id"], "used_in_planets": int(source_row is row), **values})
    summaries = []
    for s in sorted((s for s in stars if selected(s, radius, exclude_sol)), key=lambda s: (float(s["dist"]), int(s["id"]))):
        ps = by_star[s["id"]]
        solar_count = 8 if is_sol(s) else 0
        count = len(ps) + solar_count
        summary = dict.fromkeys(STAR_COLUMNS, "")
        summary.update(star_name=star_name(s), has_planets="Yes" if count else "Unknown", planet_count=count,
                       confirmed_exoplanet_count=len(ps), distance_ly=scaled(s["dist"], PC_TO_LY),
                       planet_names="; ".join(p["planet_name"] for p in ps) if not solar_count else
                       "Mercury; Venus; Earth; Mars; Jupiter; Saturn; Uranus; Neptune",
                       notes="Sol: 8 Solar System planets; outside NASA exoplanet export." if solar_count else "",
                       hyg_id=s["id"], nasa_hosts="; ".join(sorted({p["host_name"] for p in ps})),
                       match_method="; ".join(sorted({p["match_method"] for p in ps})),
                       controversial_planet_count=sum(p["controversial"] for p in ps),
                       distance_pc=scaled(s["dist"]), spectral_type=s.get("spect", ""),
                       hyg_luminosity_solar=scaled(s.get("lum")), apparent_magnitude=scaled(s.get("mag")),
                       hip=s.get("hip", ""), hd=s.get("hd", ""), hr=s.get("hr", ""), gliese=s.get("gl", ""),
                       bayer_flamsteed=s.get("bf", ""), proper_name=s.get("proper", ""),
                       ra_deg_j2000=scaled(s.get("ra"), 15), dec_deg_j2000=scaled(s.get("dec")),
                       **{f"{axis}_ly": value for axis, value in cartesian(s, PC_TO_LY).items()},
                       **{f"{axis}_pc": value for axis, value in cartesian(s).items()},
                       component=s.get("comp", ""), primary_hyg_id=s.get("comp_primary", ""),
                       archive_planet_count=len(ps), solar_system_planet_count=solar_count)
        summaries.append(summary)
    return summaries, planets, measurements


def safe_csv(value):
    # Preserve literal catalog strings when opening a CSV in spreadsheet software.
    return "'" + value if isinstance(value, str) and value.startswith(("=", "+", "-", "@")) else value


def write_csv(path, rows, columns):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=columns)
        writer.writeheader()
        writer.writerows({k: safe_csv(row.get(k, "")) for k in columns} for row in rows)


def write_workbook(path, tables):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo
        from openpyxl.worksheet.datavalidation import DataValidation
        from openpyxl.formatting.rule import FormulaRule
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise ValueError("Excel output requires: python -m pip install -r requirements.txt (or use --csv-only)") from exc
    wb = Workbook()
    wb.remove(wb.active)
    for index, (title, (rows, columns)) in enumerate(tables.items(), 1):
        ws = wb.create_sheet(title)
        ws.append(columns)
        for row in rows:
            ws.append([row.get(c, "") for c in columns])
        for row in ws:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # source text must never become a formula
                elif isinstance(cell.value, float):
                    cell.number_format = "0.######"
        ws.freeze_panes = "B2"
        ws.row_dimensions[1].height = 34
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="183C54")
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
        for n, col in enumerate(columns, 1):
            width = min(34, max(16, len(col) + 2))
            if col in ("notes", "description", "value", "planet_names", "evidence"):
                width = 65
            ws.column_dimensions[get_column_letter(n)].width = width
            if col in {"system_id", "faction_owners"}:
                ws.column_dimensions[get_column_letter(n)].width = 30
                for row_number in range(2, ws.max_row + 1):
                    ws.cell(row_number, n).fill = PatternFill("solid", fgColor="FFF2CC")
        if rows:
            tab = Table(displayName=f"Data{index}", ref=ws.dimensions)
            tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(tab)
        if title == "Stars":
            for n, row in enumerate(rows, 2):
                # Whole-column references also count new rows appended to Planets.
                ws[f"E{n}"] = f'=COUNTIFS(Planets!$C:$C,I{n},Planets!$E:$E,1)'
                ws[f"C{n}"] = f'=IF(D{n}="",E{n}+{row["solar_system_planet_count"]},D{n})'
                ws[f"B{n}"] = f'=IF(C{n}>0,"Yes",IF(D{n}="","Unknown","No"))'
                for col in ("D", "H"):
                    ws[f"{col}{n}"].fill = PatternFill("solid", fgColor="FFF2CC")
            dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1=0, allow_blank=True)
            dv.error = "Enter a nonnegative integer, or leave blank to use planet rows."
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add("D2:D1048576")
            ws.conditional_formatting.add(f"B2:C{max(2, ws.max_row)}", FormulaRule(
                formula=['$B2="Yes"'], fill=PatternFill("solid", fgColor="C6EFCE")))
        if title == "Planets":
            dv = DataValidation(type="whole", operator="between", formula1=0, formula2=1)
            dv.showErrorMessage = True
            ws.add_data_validation(dv)
            dv.add("E2:E1048576")
            for n in range(2, ws.max_row + 1):
                for col in ("E", "F"):
                    ws[f"{col}{n}"].fill = PatternFill("solid", fgColor="FFF2CC")
    wb.save(path)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--radius-ly", type=float, default=DEFAULT_RADIUS_LY)
    parser.add_argument("--hyg", type=Path, default=ROOT / "hygdata_v42.csv.gz")
    parser.add_argument("--planets", type=Path, help="Defaults to newest lexically sorted PS_2026*.csv beside script")
    parser.add_argument("--aliases", type=Path, default=ROOT / "host_aliases.csv")
    parser.add_argument("--no-aliases", action="store_true")
    parser.add_argument("--output-dir", type=Path)
    parser.add_argument("--review-arcsec", type=float, default=180, help="Position candidate radius; never assigns planets")
    parser.add_argument("--exclude-sol", action="store_true")
    parser.add_argument("--csv-only", action="store_true", help="No third-party dependencies needed")
    parser.add_argument("--campaign-from", type=Path, help="Preserve campaign/count/notes fields from this workbook; defaults to existing output workbook")
    args = parser.parse_args(argv)
    if not math.isfinite(args.radius_ly) or args.radius_ly <= 0:
        parser.error("--radius-ly must be positive and finite")
    if not math.isfinite(args.review_arcsec) or args.review_arcsec <= 0:
        parser.error("--review-arcsec must be positive and finite")
    try:
        if args.planets is None:
            files = sorted(ROOT.glob("PS_2026*.csv"))
            if not files:
                raise ValueError("No PS_2026*.csv found; supply --planets")
            args.planets = files[-1]
        stars, _ = read_csv(args.hyg, ["id", "dist", "ra", "dec"])
        raw, nasa_columns = read_csv(args.planets, ["pl_name", "hostname", "default_flag"])
        chosen, grouped = choose_planets(raw)
        alias_path = None if args.no_aliases else args.aliases
        matches, audit = match_hosts(stars, chosen, args.radius_ly, alias_path, args.review_arcsec, args.exclude_sol)
        summaries, planets, measurements = make_tables(stars, chosen, grouped, matches, args.radius_ly, args.exclude_sol)
        archive_planet_count = len(planets)
        archive_host_count = len({planet["hyg_id"] for planet in planets})
        output = args.output_dir or ROOT / "output" / f"nearby_{args.radius_ly:g}ly"
        existing = output / "nearby_stars.xlsx"
        campaign_source = args.campaign_from or (existing if existing.exists() else None)
        factions = apply_campaign(summaries, planets, campaign_source)
        warnings = defaultdict(list)
        for entry in audit:
            if entry["status"] == "matched_selected" and entry["notes"]:
                warnings[entry["hyg_id"]].append(f"{entry['nasa_host']}: {entry['notes']}")
        for summary in summaries:
            notes = summary["notes"]
            for warning in warnings[summary["hyg_id"]]:
                repeated = f"{warning} {warning}"
                while repeated in notes:
                    notes = notes.replace(repeated, warning)
                if warning not in notes:
                    notes = " ".join(filter(None, (notes, warning)))
            summary["notes"] = notes
        review = [r for r in audit if r["status"] in {"unmatched_nearby", "ambiguous_identifier"}
                  or (r["status"] == "matched_outside_radius" and r["nasa_distance_ly"] != ""
                      and r["nasa_distance_ly"] <= args.radius_ly)
                  or (r["status"] == "matched_selected" and r["notes"])]
        metadata = {
            "radius_ly": args.radius_ly, "parsec_to_light_year": PC_TO_LY,
            "hyg_file": args.hyg.name, "hyg_sha256": hashlib.sha256(args.hyg.read_bytes()).hexdigest(),
            "nasa_file": args.planets.name, "nasa_sha256": hashlib.sha256(args.planets.read_bytes()).hexdigest(),
            "aliases_file": metadata_filename(alias_path),
            "aliases_sha256": hashlib.sha256(alias_path.read_bytes()).hexdigest() if alias_path else None,
            "input_hyg_entries": len(stars), "input_nasa_rows": len(raw), "unique_archive_planets": len(chosen),
            "selected_hyg_entries": len(summaries), "matched_exoplanet_hosts": archive_host_count,
            "matched_confirmed_exoplanets": archive_planet_count,
            "preserved_campaign_planets": len(planets) - archive_planet_count, "review_rows": len(review),
            "include_sol": not args.exclude_sol, "review_arcsec": args.review_arcsec,
            "coordinate_frame": "Sol-centered J2000 equatorial; +X vernal equinox, +Y RA 6h, +Z north celestial pole",
            "campaign_source": metadata_filename(campaign_source),
        }
        instructions = [
            ("Faction ownership", OWNERSHIP_HELP),
            ("Editing", "Stars: yellow planet_count_override accepts a count; leave blank to count Planets. Yellow notes cells are free text. Planets: include_in_count=0 excludes a row. Formulas update in Excel/LibreOffice."),
            ("Planet status", "Yes means at least one counted planet. Unknown means no confirmed exoplanet was matched in these files. Explicit count override 0 displays No (your assertion). archive_planet_count preserves the original matched count."),
            ("Counts", "Count distinct pl_name values, not publication rows or sy_pnum (which is system-wide). Controversial archive members remain included and flagged. Sol has 8 known planets but 0 exoplanets; Solar System orbital data are not supplied."),
            ("Naming", "Preserve NASA scientific designations exactly. Host + existing lowercase planet letter; letters reflect discovery order, not orbital distance. Uppercase stellar components remain distinct. HYG proper names are display labels, not newly assigned IAU names."),
            ("Selection", "HYG distance in parsecs controls the inclusive radius cut; missing/nonfinite/negative/sentinel distances are excluded. HYG is not a complete census and can contain duplicates. Rows are HYG entries, not deduplicated physical stars."),
            ("Cartesian coordinates", "Stars includes x_ly/y_ly/z_ly in light-years and x_pc/y_pc/z_pc in parsecs. These are HYG heliocentric J2000 equatorial coordinates: +X toward the vernal equinox, +Y toward RA 6 hours, +Z toward the north celestial pole. Sol is pinned to exactly (0,0,0); other HYG positions are preserved."),
            ("Matching", "Exact normalized catalog, Bayer, Flamsteed, variable or proper name, then explicit verified aliases. Never assign by angular proximity alone; Review lists candidates. Matches outside the radius are retained in Host matches."),
            ("Distance checks", "NASA and HYG distances can differ substantially; >20% differences are flagged. Angular separations are diagnostic only: epochs/proper motion can differ. No positions are propagated."),
            ("Measurements", "Planets uses one default NASA row per planet. If absent, latest row is used and flagged. No values are filled from other publications. Measurements retains all published rows, errors, limits and references for matched planets; used_in_planets=1 marks chosen rows."),
            ("Habitability", "Semi-major axis is orbital size, not instantaneous separation. Published flux and equilibrium temperature are indicators, not habitability verdicts; equilibrium temperature is not surface temperature. Mass may be M sin(i), indicated by mass_provenance. Blank means unavailable. See Measurements for uncertainty and limit flags."),
            ("Luminosity", "HYG luminosity is supplied as-is; no bolometric correction or habitable-zone calculation is attempted. Age, atmosphere, rotation and activity are not reliably available in this export."),
            ("Limits", "NASA parameter lim: 0 = measured value, 1 = upper limit, -1 = lower limit. Nonzero limits are called out in Planets notes; err1/err2 are upper/lower uncertainties in Measurements."),
            ("Regeneration", "Campaign fields, notes, count overrides, planet inclusion flags and manually added planet rows are preserved from the existing workbook or --campaign-from while their host HYG row remains selected. Other values are regenerated. planet_names is an archive snapshot. CSV files are independent snapshots."),
            ("HYG source", "David Nash, HYG database, CC BY-SA 4.0: https://www.astronexus.com/projects/hyg ; field documentation: https://github.com/astronexus/HYG-Database/blob/main/hyg/README.md"),
            ("NASA source", "NASA Exoplanet Archive (NASA Exoplanet Science Institute / Caltech): https://exoplanetarchive.ipac.caltech.edu/docs/API_PS_columns.html"),
            ("Naming source", "https://iauarchive.eso.org/public/themes/naming_exoplanets/"),
        ]
        readme = [{"topic": k, "description": v} for k, v in instructions]
        readme += [{"topic": k, "description": str(v)} for k, v in metadata.items()]
        definitions = []
        with args.planets.open(encoding="utf-8-sig") as f:
            for line in f:
                if line.startswith("# COLUMN "):
                    key, desc = line[len("# COLUMN "):].split(":", 1)
                    definitions.append({"field": key.strip(), "description": desc.strip()})
        tables = {"Stars": (summaries, STAR_COLUMNS), "Planets": (planets, PLANET_COLUMNS),
                  "Factions": (factions, FACTION_COLUMNS),
                  "Read me": (readme, ["topic", "description"]), "Review": (review, AUDIT_COLUMNS),
                  "Host matches": (audit, AUDIT_COLUMNS),
                  "Measurements": (measurements, ["hyg_id", "used_in_planets", *nasa_columns]),
                  "NASA fields": (definitions, ["field", "description"])}
        output.mkdir(parents=True, exist_ok=True)
        if not args.csv_only:
            write_workbook(output / "nearby_stars.xlsx", tables)
        for title, (rows, columns) in tables.items():
            write_csv(output / f"{title.lower().replace(' ', '_')}.csv", rows, columns)
        (output / "summary.json").write_text(json.dumps(metadata, indent=2) + "\n")
        print(f"{len(summaries)} HYG entries; {metadata['matched_exoplanet_hosts']} exoplanet hosts; "
              f"{archive_planet_count} confirmed exoplanets; {metadata['preserved_campaign_planets']} campaign planets; "
              f"{len(review)} review rows.")
        print(f"Output: {output.resolve()}")
        return 0
    except (ValueError, OSError, csv.Error) as exc:
        parser.exit(1, f"Error: {exc}\n")


if __name__ == "__main__":
    sys.exit(main())
